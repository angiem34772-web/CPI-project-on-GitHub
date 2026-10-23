"""Scores the out-of-sample backtest on the workbook's Backtest tab.

Each backtest row holds the model's forecast made with only the data published
before that release, next to what BLS actually printed.
"""

from __future__ import annotations

import math
from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook

from .model import round_half_up

FIRST_ROW = 5


@dataclass
class Print:
    month: str
    period: str
    actual_nsa: float
    model_nsa: float
    actual_yoy: float | None   # missing for a few early months
    model_yoy: float | None
    actual_sa: float      # first-published SA change, already one decimal
    model_sa: float

    @property
    def nsa_error(self) -> float:
        return self.model_nsa - self.actual_nsa

    @property
    def yoy_bracket_miss(self) -> float | None:
        """How far the rounded YoY forecast missed the rounded print, in points."""
        if self.actual_yoy is None:
            return None
        return round_half_up(round_half_up(self.model_yoy) - round_half_up(self.actual_yoy))

    @property
    def sa_miss(self) -> float:
        return abs(round_half_up(round_half_up(self.model_sa) - self.actual_sa))


@dataclass
class Score:
    n: int
    mae_nsa: float
    rmse_nsa: float
    bias_nsa: float
    yoy_exact: float
    yoy_within_one: float
    sa_exact: float
    sa_miss_per_print: float


def _optional(value) -> float | None:
    return None if value in (None, "") else float(value)


def load_prints(path: str | Path) -> list[Print]:
    ws = load_workbook(path, data_only=True)["Backtest"]
    prints = []
    row = FIRST_ROW
    while ws[f"A{row}"].value:
        if ws[f"B{row}"].value is not None:
            prints.append(Print(
                month=str(ws[f"A{row}"].value),
                period=ws[f"L{row}"].value,
                actual_nsa=float(ws[f"B{row}"].value),
                model_nsa=float(ws[f"C{row}"].value),
                actual_yoy=_optional(ws[f"E{row}"].value),
                model_yoy=_optional(ws[f"F{row}"].value),
                actual_sa=float(ws[f"H{row}"].value),
                model_sa=float(ws[f"I{row}"].value),
            ))
        row += 1
    return prints


def score(prints: list[Print]) -> Score:
    n = len(prints)
    errors = [p.nsa_error for p in prints]
    yoy_misses = [p.yoy_bracket_miss for p in prints if p.yoy_bracket_miss is not None]
    return Score(
        n=n,
        mae_nsa=sum(abs(e) for e in errors) / n,
        rmse_nsa=math.sqrt(sum(e * e for e in errors) / n),
        bias_nsa=sum(errors) / n,
        yoy_exact=sum(m == 0 for m in yoy_misses) / len(yoy_misses),
        yoy_within_one=sum(abs(m) <= 0.1001 for m in yoy_misses) / len(yoy_misses),
        sa_exact=sum(p.sa_miss == 0 for p in prints) / n,
        sa_miss_per_print=sum(p.sa_miss for p in prints) / n,
    )


def score_by_period(prints: list[Print]) -> dict[str, Score]:
    periods: dict[str, list[Print]] = {}
    for p in prints:
        periods.setdefault(p.period, []).append(p)
    return {name: score(group) for name, group in periods.items()}
